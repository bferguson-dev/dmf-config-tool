"""Workbook builder test helper."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from dmf_tool.core.parser import METADATA_HEADERS, SHEET_HEADERS


class WorkbookBuilder:
    """Programmatically build parser test workbooks."""

    def __init__(self) -> None:
        self._workbook = Workbook()
        default_sheet = self._workbook.active
        assert default_sheet is not None
        self._workbook.remove(default_sheet)
        self._seed_metadata()
        self._seed_table_sheets()

    def _seed_metadata(self) -> None:
        worksheet = self._workbook.create_sheet("metadata")
        worksheet.append(METADATA_HEADERS)
        worksheet.append(("dmf_version", "8.8"))
        worksheet.append(("workbook_schema_version", "1.0"))

    def _seed_table_sheets(self) -> None:
        sample_rows: dict[str, tuple[object, ...]] = {
            "controllers": (
                "controller-1",
                "10.0.0.10",
                "cluster-a",
                "single",
                "",
                "site-a",
                "",
            ),
            "switches": ("leaf-1", "10.0.0.1", "filter", "site-a", "7050", ""),
            "interfaces": (
                "leaf-1",
                "ethernet1",
                "filter",
                "",
                "10g",
                "100",
                "yes",
            ),
            "groups": ("group-a", "filter", ""),
            "group_members": ("group-a", "leaf-1", "ethernet1"),
            "policies": (
                "policy-a",
                "100",
                "leaf-1:ethernet1",
                "",
                "",
                "",
                "forward",
                "yes",
                "",
            ),
            "match_rules": (
                "policy-a",
                "10",
                "10.0.0.1",
                "10.0.0.2",
                "80",
                "443",
                "100",
                "tcp",
            ),
            "service_nodes": ("service-1", "leaf-1", "ethernet1", "10.0.0.20", ""),
            "analytics_nodes": (
                "analytics-1",
                "leaf-1",
                "ethernet1",
                "10.0.0.21",
                "",
            ),
            "recorder_nodes": (
                "recorder-1",
                "leaf-1",
                "ethernet1",
                "10.0.0.22",
                "",
            ),
            "fabric_settings": (
                "fabric-a",
                "10.0.0.30",
                "10.0.0.31",
                "public",
                "enable",
                "token",
            ),
            "sites": ("site-a", "Dallas", "America/Chicago", ""),
        }
        for sheet_name, headers in SHEET_HEADERS.items():
            worksheet = self._workbook.create_sheet(sheet_name)
            worksheet.append(headers)
            worksheet.append(sample_rows[sheet_name])

    def with_missing_tab(self, tab_name: str) -> WorkbookBuilder:
        """Remove a sheet from the workbook."""
        if tab_name in self._workbook.sheetnames:
            self._workbook.remove(self._workbook[tab_name])
        return self

    def with_invalid_ip(
        self,
        tab: str,
        row: int,
        col: str,
        value: str,
    ) -> WorkbookBuilder:
        """Assign a syntactically invalid IP value to a cell."""
        self._workbook[tab][f"{col}{row}"] = value
        return self

    def with_merged_cells(self, tab: str, cell_range: str) -> WorkbookBuilder:
        """Create a merged cell range."""
        self._workbook[tab].merge_cells(cell_range)
        return self

    def with_formula(self, tab: str, row: int, col: str) -> WorkbookBuilder:
        """Insert a formula string into a cell."""
        self._workbook[tab][f"{col}{row}"] = "=SUM(1,1)"
        return self

    def with_extra_column(self, tab: str, col_name: str) -> WorkbookBuilder:
        """Append an extra column header to a tab."""
        worksheet = self._workbook[tab]
        header_column = worksheet.max_column + 1
        worksheet.cell(row=1, column=header_column, value=col_name)
        worksheet.cell(row=2, column=header_column, value="extra")
        return self

    def with_duplicate_name(self, tab: str, name: str) -> WorkbookBuilder:
        """Append a duplicate-name row for tables that carry a name column."""
        worksheet = self._workbook[tab]
        headers = [cell.value for cell in worksheet[1]]
        source_row = [
            worksheet.cell(row=2, column=index + 1).value
            for index in range(len(headers))
        ]
        duplicate_row = list(source_row)
        if "name" in headers:
            duplicate_row[headers.index("name")] = name
        if "group_name" in headers:
            duplicate_row[headers.index("group_name")] = name
        if "policy_name" in headers:
            duplicate_row[headers.index("policy_name")] = name
        worksheet.append(tuple(duplicate_row))
        return self

    def build(self) -> Workbook:
        """Return the constructed workbook."""
        return self._workbook

    def save(self, path: Path) -> Path:
        """Save the workbook to disk."""
        path.parent.mkdir(parents=True, exist_ok=True)
        self._workbook.save(path)
        return path
