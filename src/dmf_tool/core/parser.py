"""Workbook acceptance checks and parsing."""

from __future__ import annotations

import mimetypes
import zipfile
from collections import defaultdict
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from dmf_tool.core.models.workbook import (
    RawAnalyticsNode,
    RawController,
    RawFabricSettings,
    RawGroup,
    RawGroupMember,
    RawInterface,
    RawMatchRule,
    RawPolicy,
    RawRecorderNode,
    RawServiceNode,
    RawSite,
    RawSwitch,
    RawWorkbook,
)

MAX_WORKBOOK_BYTES = 10 * 1024 * 1024
XLSX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
METADATA_SHEET = "metadata"
METADATA_HEADERS = ("key", "value")

SHEET_HEADERS: dict[str, tuple[str, ...]] = {
    "controllers": (
        "name",
        "management_ip",
        "cluster_name",
        "role",
        "ha_peer",
        "site_name",
        "description",
    ),
    "switches": (
        "name",
        "management_ip",
        "role",
        "site_name",
        "model",
        "description",
    ),
    "interfaces": (
        "switch_name",
        "name",
        "role",
        "description",
        "speed",
        "vlan_id",
        "enabled",
    ),
    "groups": ("name", "role", "description"),
    "group_members": ("group_name", "switch_name", "interface_name"),
    "policies": (
        "name",
        "priority",
        "source_interface",
        "source_group",
        "delivery_interface",
        "delivery_group",
        "action",
        "enabled",
        "description",
    ),
    "match_rules": (
        "policy_name",
        "sequence",
        "source_ip",
        "destination_ip",
        "source_port",
        "destination_port",
        "vlan_id",
        "protocol",
    ),
    "service_nodes": (
        "name",
        "switch_name",
        "interface_name",
        "management_ip",
        "description",
    ),
    "analytics_nodes": (
        "name",
        "switch_name",
        "interface_name",
        "management_ip",
        "description",
    ),
    "recorder_nodes": (
        "name",
        "switch_name",
        "interface_name",
        "management_ip",
        "description",
    ),
    "fabric_settings": (
        "fabric_name",
        "syslog_server",
        "ntp_server",
        "snmp_community",
        "enable_password",
        "api_token",
    ),
    "sites": ("name", "location", "timezone", "description"),
}


@dataclass(slots=True)
class WorkbookAcceptanceError(Exception):
    """Raised when a workbook should be rejected before parsing."""

    message: str

    def __str__(self) -> str:
        return self.message


@dataclass(slots=True)
class WorkbookParseError(Exception):
    """Raised when a workbook cannot be parsed into raw models."""

    message: str

    def __str__(self) -> str:
        return self.message


def _coerce_cell_value(value: object) -> str | None:
    """Convert a workbook cell value to a permissive string field."""
    if value is None:
        return None
    return str(value)


def _sheet_or_none(workbook: Workbook, title: str) -> Worksheet | None:
    """Return a worksheet when it exists."""
    return workbook[title] if title in workbook.sheetnames else None


def _iter_data_rows(worksheet: Worksheet) -> Iterable[dict[str, str | None]]:
    """Yield row dictionaries using the sheet header row."""
    headers = [_coerce_cell_value(cell.value) or "" for cell in worksheet[1]]
    for row in worksheet.iter_rows(min_row=2):
        values = [_coerce_cell_value(cell.value) for cell in row[: len(headers)]]
        if all(value is None for value in values):
            continue
        yield dict(zip(headers, values, strict=False))


def _validate_file_extension(path: Path) -> None:
    """Validate supported workbook extensions and mime type."""
    suffix = path.suffix.lower()
    if suffix == ".xlsm":
        raise WorkbookAcceptanceError("Macro-enabled .xlsm workbooks are unsupported")
    if suffix == ".xls":
        raise WorkbookAcceptanceError("Legacy .xls workbooks are unsupported")
    if suffix != ".xlsx":
        raise WorkbookAcceptanceError("Workbook must use the .xlsx extension")

    mime_type, _ = mimetypes.guess_type(path.name)
    if mime_type != XLSX_MIME_TYPE:
        raise WorkbookAcceptanceError("Workbook MIME type must be .xlsx")


def _validate_zip_structure(path: Path) -> None:
    """Reject macro-enabled or externally linked workbooks before parsing."""
    try:
        with zipfile.ZipFile(path) as workbook_zip:
            members = set(workbook_zip.namelist())
            if "xl/vbaProject.bin" in members:
                raise WorkbookAcceptanceError("Workbook contains VBA macros")
            if any(member.startswith("xl/externalLinks/") for member in members):
                raise WorkbookAcceptanceError("Workbook contains external links")
    except zipfile.BadZipFile as exc:
        raise WorkbookAcceptanceError("Workbook is corrupt or unreadable") from exc


def validate_workbook_path(workbook_path: Path) -> None:
    """Run pre-parse acceptance checks against a workbook path."""
    if not workbook_path.exists():
        raise WorkbookAcceptanceError("Workbook path does not exist")
    if not workbook_path.is_file():
        raise WorkbookAcceptanceError("Workbook path must be a file")
    if workbook_path.stat().st_size > MAX_WORKBOOK_BYTES:
        raise WorkbookAcceptanceError("Workbook exceeds the 10MB size limit")

    _validate_file_extension(workbook_path)
    _validate_zip_structure(workbook_path)


def _detect_unsupported_cells(workbook: Workbook) -> None:
    """Reject hidden sheets and formula cells."""
    for worksheet in workbook.worksheets:
        if worksheet.sheet_state != "visible":
            raise WorkbookAcceptanceError(
                f"Workbook contains hidden sheet: {worksheet.title}"
            )
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    raise WorkbookAcceptanceError(
                        "Workbook contains formula cell: "
                        f"{worksheet.title}!{cell.coordinate}"
                    )
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    raise WorkbookAcceptanceError(
                        "Workbook contains formula-like cell: "
                        f"{worksheet.title}!{cell.coordinate}"
                    )


def _load_workbook(path: Path) -> Workbook:
    """Load a workbook after file-level checks."""
    try:
        workbook = load_workbook(path, read_only=False, keep_vba=False, data_only=False)
    except (InvalidFileException, OSError, ValueError, zipfile.BadZipFile) as exc:
        raise WorkbookAcceptanceError("Workbook is corrupt or unreadable") from exc

    if getattr(workbook, "vba_archive", None) is not None:
        raise WorkbookAcceptanceError("Workbook contains VBA macros")

    _detect_unsupported_cells(workbook)
    return workbook


def _parse_metadata(workbook: Workbook) -> tuple[str, str]:
    """Read workbook metadata from the metadata sheet."""
    worksheet = _sheet_or_none(workbook, METADATA_SHEET)
    if worksheet is None:
        return "", ""

    metadata: dict[str, str] = {}
    for row in _iter_data_rows(worksheet):
        key = (row.get("key") or "").strip()
        if not key:
            continue
        metadata[key] = row.get("value") or ""
    return metadata.get("dmf_version", ""), metadata.get("workbook_schema_version", "")


def _parse_simple_rows[T](
    workbook: Workbook,
    sheet_name: str,
    model_type: type[T],
) -> list[T]:
    """Parse a simple sheet directly into raw row models."""
    worksheet = _sheet_or_none(workbook, sheet_name)
    if worksheet is None:
        return []
    return [model_type(**row) for row in _iter_data_rows(worksheet)]


def _parse_groups(workbook: Workbook) -> list[RawGroup]:
    """Parse groups plus nested group members."""
    group_members_sheet = _sheet_or_none(workbook, "group_members")
    members_by_group: dict[str, list[RawGroupMember]] = defaultdict(list)
    if group_members_sheet is not None:
        for row in _iter_data_rows(group_members_sheet):
            group_name = row.get("group_name") or ""
            members_by_group[group_name].append(
                RawGroupMember(
                    switch_name=row.get("switch_name"),
                    interface_name=row.get("interface_name"),
                )
            )

    groups_sheet = _sheet_or_none(workbook, "groups")
    if groups_sheet is None:
        return []

    groups: list[RawGroup] = []
    for row in _iter_data_rows(groups_sheet):
        group_name = row.get("name") or ""
        groups.append(
            RawGroup(
                name=row.get("name"),
                role=row.get("role"),
                description=row.get("description"),
                members=members_by_group.get(group_name, []),
            )
        )
    return groups


def _parse_policies(workbook: Workbook) -> list[RawPolicy]:
    """Parse policies plus nested match rules."""
    rules_sheet = _sheet_or_none(workbook, "match_rules")
    rules_by_policy: dict[str, list[RawMatchRule]] = defaultdict(list)
    if rules_sheet is not None:
        for row in _iter_data_rows(rules_sheet):
            policy_name = row.get("policy_name") or ""
            rules_by_policy[policy_name].append(
                RawMatchRule(
                    sequence=row.get("sequence"),
                    source_ip=row.get("source_ip"),
                    destination_ip=row.get("destination_ip"),
                    source_port=row.get("source_port"),
                    destination_port=row.get("destination_port"),
                    vlan_id=row.get("vlan_id"),
                    protocol=row.get("protocol"),
                )
            )

    policies_sheet = _sheet_or_none(workbook, "policies")
    if policies_sheet is None:
        return []

    policies: list[RawPolicy] = []
    for row in _iter_data_rows(policies_sheet):
        policy_name = row.get("name") or ""
        policies.append(
            RawPolicy(
                name=row.get("name"),
                priority=row.get("priority"),
                source_interface=row.get("source_interface"),
                source_group=row.get("source_group"),
                delivery_interface=row.get("delivery_interface"),
                delivery_group=row.get("delivery_group"),
                action=row.get("action"),
                enabled=row.get("enabled"),
                description=row.get("description"),
                match_rules=rules_by_policy.get(policy_name, []),
            )
        )
    return policies


def _parse_fabric_settings(workbook: Workbook) -> RawFabricSettings | None:
    """Parse a single-row fabric settings sheet."""
    sheet = _sheet_or_none(workbook, "fabric_settings")
    if sheet is None:
        return None
    rows = list(_iter_data_rows(sheet))
    if not rows:
        return None
    return RawFabricSettings(**rows[0])


def parse_workbook(workbook_path: Path) -> RawWorkbook:
    """Validate and parse a workbook into raw models."""
    validate_workbook_path(workbook_path)
    workbook = _load_workbook(workbook_path)
    try:
        dmf_version, workbook_schema_version = _parse_metadata(workbook)
        return RawWorkbook(
            dmf_version=dmf_version,
            workbook_schema_version=workbook_schema_version,
            controllers=_parse_simple_rows(workbook, "controllers", RawController),
            switches=_parse_simple_rows(workbook, "switches", RawSwitch),
            interfaces=_parse_simple_rows(workbook, "interfaces", RawInterface),
            groups=_parse_groups(workbook),
            policies=_parse_policies(workbook),
            service_nodes=_parse_simple_rows(
                workbook,
                "service_nodes",
                RawServiceNode,
            ),
            analytics_nodes=_parse_simple_rows(
                workbook,
                "analytics_nodes",
                RawAnalyticsNode,
            ),
            recorder_nodes=_parse_simple_rows(
                workbook,
                "recorder_nodes",
                RawRecorderNode,
            ),
            fabric_settings=_parse_fabric_settings(workbook),
            sites=_parse_simple_rows(workbook, "sites", RawSite),
        )
    except WorkbookAcceptanceError:
        raise
    except Exception as exc:
        raise WorkbookParseError("Failed to parse workbook contents") from exc
